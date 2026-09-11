from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "ANDROID_TESTER_INSTRUCTION_RU.md"
OUTPUT = ROOT / "LectureVault-Android-инструкция-для-тестировщиков.docx"

BLACK = "000000"
MUTED = "5E6470"
ACCENT = "E85D2A"
PALE = "F7F7F8"
LINE = "D9D9D9"


def set_font(run, name="Aptos", size=None, bold=None, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(run, fill="ECEEF1"):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    run._element.get_or_add_rPr().append(shd)


def hyperlink(paragraph, url, label=None):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    node = OxmlElement("w:hyperlink")
    node.set(qn("r:id"), rel_id)
    run_node = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), ACCENT)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([color, underline])
    run_node.append(props)
    text = OxmlElement("w:t")
    text.text = label or url
    run_node.append(text)
    node.append(run_node)
    paragraph._p.append(node)


TOKEN = re.compile(r"(`[^`]+`|https?://\S+)")


def add_inline(paragraph, text):
    pos = 0
    for match in TOKEN.finditer(text):
        if match.start() > pos:
            set_font(paragraph.add_run(text[pos:match.start()]), size=10.5)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, name="Cascadia Mono", size=9.5, color="24262B")
            shade(run)
        else:
            hyperlink(paragraph, token.rstrip(".,)"))
            tail = token[len(token.rstrip(".,)")):]
            if tail:
                set_font(paragraph.add_run(tail), size=10.5)
        pos = match.end()
    if pos < len(text):
        set_font(paragraph.add_run(text[pos:]), size=10.5)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_fill(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def set_cell_borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), LINE)
        borders.append(element)


def add_modes_table(doc):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Cm(3.2), Cm(6.2), Cm(6.2))
    headers = ("Режим", "Что требуется", "Когда использовать")
    for index, cell in enumerate(table.rows[0].cells):
        cell.width = widths[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_fill(cell, "30343B")
        set_cell_borders(cell)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        run = paragraph.add_run(headers[index])
        set_font(run, size=9.5, bold=True, color="FFFFFF")
    rows = (
        ("Groq и Gemini", "Интернет, два личных API ключа и согласие на отправку данных", "Лучшее качество расшифровки и структурированного конспекта"),
        ("Локальная модель", "Около 500 МБ памяти и однократная загрузка модели", "Работа без API и резервный режим при проблемах с облаком"),
    )
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].width = widths[index]
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_fill(cells[index], "FFFFFF" if row_index % 2 == 0 else PALE)
            set_cell_borders(cells[index])
            set_cell_margins(cells[index])
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, value)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def footer_with_page_number(section):
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    run = paragraph.add_run("LectureVault 1.3.0   •   ")
    set_font(run, size=8, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def remove_paragraph_borders(style):
    p_pr = style.element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    footer_with_page_number(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size, before, after in (
        ("Title", 27, 0, 8),
        ("Heading 1", 17, 15, 6),
        ("Heading 2", 13, 11, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display" if style_name != "Heading 2" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        remove_paragraph_borders(style)

    title_done = False
    inserted_modes = False
    intro_done = False
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# ") and not title_done:
            p = doc.add_paragraph(style="Title")
            p.add_run("LectureVault для Android")
            title_done = True
            subtitle = doc.add_paragraph()
            subtitle.paragraph_format.space_after = Pt(14)
            run = subtitle.add_run("Инструкция по установке и настройке для тестировщиков")
            set_font(run, size=12, color=MUTED)
            continue
        if line == "Эту инструкцию можно целиком переслать тестировщикам.":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(14)
            run = p.add_run("Документ объясняет назначение приложения, установку, подключение Obsidian, варианты обработки и способы устранения основных ошибок.")
            set_font(run, size=10.5, color=MUTED)
            intro_done = True
            continue
        if line.startswith("## "):
            heading = line[3:].rstrip(".:!?")
            doc.add_paragraph(heading, style="Heading 1")
            if heading.endswith("Выбор обработки") and not inserted_modes:
                add_modes_table(doc)
                inserted_modes = True
            continue
        if line.startswith("### "):
            doc.add_paragraph(line[4:].rstrip(".:!?"), style="Heading 2")
            continue
        bullet = re.match(r"^-\s+(.*)$", line)
        numbered = re.match(r"^(\d+)\.\s+(.*)$", line)
        if bullet:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.55)
            p.paragraph_format.first_line_indent = Cm(-0.35)
            p.paragraph_format.space_after = Pt(3)
            marker = p.add_run("•  ")
            set_font(marker, size=11, bold=True, color=ACCENT)
            add_inline(p, bullet.group(1))
        elif numbered:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.72)
            p.paragraph_format.first_line_indent = Cm(-0.52)
            p.paragraph_format.space_after = Pt(4)
            marker = p.add_run(f"{numbered.group(1)}.  ")
            set_font(marker, size=10.5, bold=True, color=ACCENT)
            add_inline(p, numbered.group(2))
        else:
            p = doc.add_paragraph()
            p.paragraph_format.widow_control = True
            add_inline(p, line)

    props = doc.core_properties
    props.title = "LectureVault для Android"
    props.subject = "Инструкция по установке и настройке для тестировщиков"
    props.author = "LectureVault"
    props.keywords = "LectureVault, Android, Obsidian, Groq, Gemini"
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
